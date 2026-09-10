"""Read-only delivery QA; workbook authoring is performed by Artifact Tool."""
import json
from collections import Counter
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

path = Path("outputs/rahoitusrekisteri.xlsx")
expected = json.loads(Path("tmp/workbook-data.json").read_text())
counts = Counter(r["effective_status"] for r in expected["rows"])
total, source_count, funder_count = len(expected["rows"]), len(expected["sources"]), len(expected["funders"])
book = load_workbook(path, data_only=True, read_only=False)
assert book.sheetnames == ["Aloita tästä", "Haut", "Lähteet", "Rahoittajat"]
summary = book["Aloita tästä"]
assert [summary[f"B{i}"].value for i in range(6, 13)] == [total]+[counts[s] for s in ["open", "forthcoming", "rolling", "closed", "unknown", "cancelled"]]
assert summary["B13"].value == sum(r["record_kind"] == "funding_scheme" for r in expected["rows"])
assert summary["B14"].value == sum(r["record_kind"] == "advance_information" for r in expected["rows"])
for index, priority in enumerate(expected["priorities"], 41):
    assert summary[f"B{index}"].value == priority["records"], priority
assert book["Haut"].max_row == total+5
assert book["Lähteet"].max_row == source_count+5
assert isinstance(book["Haut"]["D6"].value, datetime)
assert isinstance(book["Lähteet"]["D6"].value, datetime)
assert isinstance(summary["B28"].value, datetime)
errors = [(sheet.title, cell.coordinate, cell.value) for sheet in book for row in sheet for cell in row if cell.data_type == "e"]
assert not errors, errors[:10]
ids = [row[0] for row in book["Haut"].iter_rows(min_row=6, min_col=12, max_col=12, values_only=True)]
assert set(ids) == {r["id"] for r in expected["rows"]}
funder_lookup = {r["funder"]: r for r in expected["funders"]}
for row in book["Rahoittajat"].iter_rows(min_row=6, max_row=funder_count+5, values_only=True):
    original = funder_lookup[row[0]]
    assert list(row[1:6]) == [original[k] for k in ["records", "calls", "schemes", "open_or_forthcoming", "unknown_status"]], row
    assert row[7] == original["advance_information"], row
states = Counter(r["content_state"] for r in expected["rows"])
assert [summary[f"B{i}"].value for i in (31,32,33)] == [states[k] for k in ("collected_unverified","partial","in_progress")]
lookup = {r["id"]:r for r in expected["rows"]}
for row in book["Haut"].iter_rows(min_row=6, values_only=True):
    item=lookup[row[11]]
    assert list(row[17:22]) == [item[k] for k in ("document_count","fetched_count","failed_document_count","document_text_issues","pending_document_count")]
    assert row[23] == ("Kyllä" if item["limited_document_count"] else "Ei")
    names = {p["id"]: f'{p["id"]:02d} {p["name"]}' for p in expected["priorities"]}
    assert row[24] == (" · ".join(names[i] for i in item["priorities"]) or "Muut lähteet")
    assert row[25] == ("Kyllä" if item["cascade"] else "Ei")
book.close()
structure = load_workbook(path, data_only=False)
for name, ref in [("Haut", f"A5:Z{total+5}"), ("Lähteet", f"A5:J{source_count+5}"), ("Rahoittajat", f"A5:H{funder_count+5}")]:
    sheet = structure[name]
    assert sheet.freeze_panes == "B6"
    table = list(sheet.tables.values())[0]
    assert table.ref == ref and table.autoFilter.ref == ref
assert structure["Haut"]["K6"].value.startswith('=IFERROR(HYPERLINK("https://')
for row in structure["Haut"].iter_rows(min_row=6):
    relative = f"data/documents/calls/{row[11].value}.html"
    assert relative in row[22].value and (path.parent/relative).is_file()
structure.close()
print(json.dumps({"xlsx": "passed", "records": len(ids), "sources": source_count, "funders": funder_count, "formula_errors": 0, "filters": True, "dates_typed": True, "freeze_panes": "B6"}))
